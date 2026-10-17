"""Parses a pharmacy's stock sheet (CSV or XLSX) into normalised rows."""
import csv
import io
import re
from decimal import Decimal, InvalidOperation

HEADERS = {
    "medicine": {"medicine", "name", "product", "drug", "item", "medicine_name", "product_name"},
    "barcode": {"barcode", "gtin", "ean", "upc", "code"},
    "price": {"price", "unit_price", "cost"},
    "status": {"status", "availability", "availability_status", "stock_status"},
    "stock": {"stock", "quantity", "qty", "stock_quantity", "on_hand"},
}
MAX_ROWS = 5000


class SheetError(ValueError):
    pass


def _key(h) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(h or "").strip().lower()).strip("_")


def _map_headers(raw: list) -> dict[int, str]:
    out = {}
    for i, h in enumerate(raw):
        for field, names in HEADERS.items():
            if _key(h) in names and field not in out.values():
                out[i] = field
    return out


def _read_table(filename: str, data: bytes) -> list[list]:
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        try:
            ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
            return [list(r) for r in ws.iter_rows(values_only=True)]
        except Exception as e:  # corrupt / not a real xlsx
            raise SheetError("Could not read the Excel file.") from e
    if name.endswith((".csv", ".txt", ".tsv")) or not name:
        text = data.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return list(csv.reader(io.StringIO(text), dialect))
    raise SheetError("Upload a .csv or .xlsx file.")


def parse_price(v):
    if v is None or str(v).strip() == "":
        return None
    s = str(v).strip().replace(" ", "")
    if re.fullmatch(r"\d{1,3}(\.\d{3})+,\d+", s):      # 1.234,56
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(,\d{3})+(\.\d+)?", s):  # 1,234.56
        s = s.replace(",", "")
    else:
        s = s.replace(",", ".")
    s = re.sub(r"[^\d.]", "", s)
    try:
        d = Decimal(s)
    except InvalidOperation:
        raise ValueError(f"bad price '{v}'")
    if d < 0:
        raise ValueError("negative price")
    return d


def parse_stock(v):
    if v is None or str(v).strip() == "":
        return None
    try:
        n = int(float(str(v).strip().replace(",", "")))
    except ValueError:
        raise ValueError(f"bad stock '{v}'")
    if n < 0:
        raise ValueError("negative stock")
    return n


def parse_status(v, stock):
    s = str(v or "").strip().lower()
    if s in ("available", "in stock", "instock", "yes", "y", "1", "true", "in_stock"):
        return "available"
    if s in ("limited", "limited_stock", "limited stock", "low", "low stock"):
        return "limited_stock"
    if s in ("out", "out of stock", "out_of_stock", "outofstock", "no", "n", "0", "false", "sold out"):
        return "out_of_stock"
    if s and s not in ("unknown", "?"):
        raise ValueError(f"unknown status '{v}'")
    if stock is None:
        return "unknown"
    return "out_of_stock" if stock == 0 else "limited_stock" if stock <= 5 else "available"


def parse_sheet(filename: str, data: bytes) -> list[dict]:
    """Returns [{row, medicine, barcode, price, status, stock, error}]."""
    table = [r for r in _read_table(filename, data) if any(str(c or "").strip() for c in r)]
    if not table:
        raise SheetError("The file is empty.")
    cols = _map_headers(table[0])
    fields = set(cols.values())
    if not ({"medicine", "barcode"} & fields) or not ({"price", "status", "stock"} & fields):
        raise SheetError(
            "Need a 'medicine' (or 'barcode') column and at least one of 'price', 'status', 'stock'.")
    if len(table) - 1 > MAX_ROWS:
        raise SheetError(f"Too many rows (max {MAX_ROWS}).")
    rows = []
    for n, raw in enumerate(table[1:], start=2):
        rec = {"row": n, "medicine": None, "barcode": None, "price": None, "status": None, "stock": None, "error": None}
        vals = {f: (raw[i] if i < len(raw) else None) for i, f in cols.items()}
        for f in ("medicine", "barcode"):
            v = str(vals.get(f) or "").strip()
            rec[f] = v or None
        try:
            if not (rec["medicine"] or rec["barcode"]):
                raise ValueError("no medicine name or barcode")
            rec["price"] = parse_price(vals.get("price"))
            rec["stock"] = parse_stock(vals.get("stock"))
            rec["status"] = parse_status(vals.get("status"), rec["stock"]) if ({"status", "stock"} & fields) else None
        except ValueError as e:
            rec["error"] = str(e)
        rows.append(rec)
    return rows
